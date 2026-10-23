"""Export all business data to one Excel workbook (one sheet per table)."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from banking.models import BankAccount
from inventory.models import Item
from ledger.models import BankTransaction
from parties.models import Party
from transactions.models import PaymentIn, PaymentOut, Purchase, PurchaseItem, Sale, SaleItem

HEADER_FILL = PatternFill('solid', fgColor='1D4ED8')
HEADER_FONT = Font(bold=True, color='FFFFFF')
MONEY_FORMAT = '#,##0.00'
DATE_FORMAT = 'dd/mm/yyyy'


def _sheet(workbook, title, headers, rows, money_cols=(), date_cols=(), widths=None):
    ws = workbook.create_sheet(title=title[:31])
    ws.append([h for h in headers])
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical='center')
    count = 0
    for row in rows:
        ws.append(list(row))
        count += 1
    for index, header in enumerate(headers, start=1):
        letter = get_column_letter(index)
        ws.column_dimensions[letter].width = (widths or {}).get(header, max(12, min(40, len(header) + 4)))
        if header in money_cols or header in date_cols:
            fmt = MONEY_FORMAT if header in money_cols else DATE_FORMAT
            for (cell,) in ws.iter_rows(min_row=2, max_row=count + 1, min_col=index, max_col=index):
                cell.number_format = fmt
    ws.freeze_panes = 'A2'
    if count:
        ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{count + 1}'
    return ws


def _f(value):
    return float(value) if value is not None else None


def export_workbook(business):
    wb = Workbook()
    wb.remove(wb.active)

    _sheet(wb, 'Parties', ['Name', 'Type', 'Mobile', 'Alternate Mobile', 'Email', 'Address', 'City', 'State',
                           'Pincode', 'GSTIN', 'Opening Balance', 'Opening Type', 'Balance', 'Receivable', 'Payable',
                           'Last Transaction', 'Notes'],
           ((p.name, p.get_party_type_display(), p.phone, p.phone_alt, p.email, p.address, p.city, p.state, p.pincode,
             p.gstin, _f(p.opening_balance), p.get_opening_balance_type_display(), _f(p.balance), _f(p.receivable),
             _f(p.payable), p.last_txn_date, p.notes)
            for p in Party.objects.filter(business=business).order_by('name')),
           money_cols={'Opening Balance', 'Balance', 'Receivable', 'Payable'}, date_cols={'Last Transaction'},
           widths={'Name': 28, 'Address': 36, 'Notes': 30})

    _sheet(wb, 'Items', ['Name', 'Code', 'Category', 'Unit', 'HSN', 'Sale Price', 'Purchase Price', 'Tax %',
                         'Opening Stock', 'Purchased', 'Sold', 'Current Stock', 'Low Stock Limit', 'Description'],
           ((i.name, i.code, i.category.name if i.category else '', i.unit, i.hsn, _f(i.sale_price),
             _f(i.purchase_price), _f(i.tax_rate), _f(i.opening_stock), _f(i.purchased_qty), _f(i.sold_qty),
             _f(i.current_stock), _f(i.low_stock_limit), i.description)
            for i in Item.objects.filter(business=business).select_related('category').order_by('name')),
           money_cols={'Sale Price', 'Purchase Price'}, widths={'Name': 30, 'Description': 30})

    for model, line_model, fk, title in ((Sale, SaleItem, 'sale', 'Sales'), (Purchase, PurchaseItem, 'purchase',
                                                                               'Purchases')):
        _sheet(wb, title, ['Date', 'Number', 'Party', 'Mobile', 'Mode', 'Subtotal', 'Discount', 'Tax', 'Charges',
                           'Round Off', 'Total', 'Paid', 'Pending', 'Status', 'Due Date', 'Notes'],
               ((b.date, b.number, b.party_name, b.party_phone, b.get_mode_display(), _f(b.subtotal),
                 _f(b.item_discount + b.discount_amount), _f(b.tax_total), _f(b.charges_amount), _f(b.round_off),
                 _f(b.total), _f(b.paid_amount), _f(b.balance_due), b.get_payment_status_display(), b.due_date, b.notes)
                for b in model.objects.filter(business=business).order_by('date', 'seq', 'id').iterator()),
               money_cols={'Subtotal', 'Discount', 'Tax', 'Charges', 'Round Off', 'Total', 'Paid', 'Pending'},
               date_cols={'Date', 'Due Date'}, widths={'Party': 26, 'Notes': 30})
        _sheet(wb, f'{title} Items', ['Date', 'Number', 'Party', 'Item', 'Code', 'Qty', 'Unit', 'Rate', 'Discount',
                                      'Tax %', 'Tax', 'Amount'],
               ((line.parent.date, line.parent.number, line.parent.party_name, line.name, line.code,
                 _f(line.quantity), line.unit, _f(line.rate), _f(line.discount_amount), _f(line.tax_rate),
                 _f(line.tax_amount), _f(line.amount))
                for line in _lines(line_model, fk, business)),
               money_cols={'Rate', 'Discount', 'Tax', 'Amount'}, date_cols={'Date'}, widths={'Item': 30, 'Party': 24})

    for model, title, bill_attr in ((PaymentIn, 'Payments In', 'sale'), (PaymentOut, 'Payments Out', 'purchase')):
        _sheet(wb, title, ['Date', 'Number', 'Party', 'Account', 'Amount', 'Reference No.', 'Against Bills',
                           'Unlinked Amount', 'Notes'],
               ((p.date, p.number, p.party_name, p.account.name, _f(p.amount), p.reference_no,
                 ', '.join(getattr(a, bill_attr).number for a in p.allocations.all()), _f(p.unallocated_amount),
                 p.notes)
                for p in model.objects.filter(business=business).select_related('account')
                .prefetch_related(f'allocations__{bill_attr}').order_by('date', 'seq', 'id')),
               money_cols={'Amount', 'Unlinked Amount'}, date_cols={'Date'}, widths={'Party': 26, 'Notes': 30})

    _sheet(wb, 'Bank Accounts', ['Name', 'Type', 'Bank Name', 'Account Name', 'Account Number', 'IFSC', 'Branch',
                                 'UPI ID', 'Opening Balance', 'Current Balance', 'Notes'],
           ((a.name, a.get_kind_display(), a.bank_name, a.account_name, a.account_number, a.ifsc, a.branch, a.upi_id,
             _f(a.opening_balance), _f(a.current_balance), a.notes)
            for a in BankAccount.objects.filter(business=business).order_by('kind', 'name')),
           money_cols={'Opening Balance', 'Current Balance'}, widths={'Name': 26})

    _sheet(wb, 'Cash & Bank Entries', ['Date', 'Account', 'Type', 'Number', 'Description', 'Reference No.',
                                       'Money In', 'Money Out'],
           ((t.date, t.account.name, t.get_txn_type_display(), t.number, t.description, t.reference_no,
             _f(t.amount_in), _f(t.amount_out))
            for t in BankTransaction.objects.filter(business=business).select_related('account')
            .order_by('date', 'source_created_at', 'id').iterator()),
           money_cols={'Money In', 'Money Out'}, date_cols={'Date'}, widths={'Description': 40, 'Account': 22})

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _lines(line_model, fk, business):
    for line in line_model.objects.filter(**{f'{fk}__business': business}).select_related(fk).order_by(
            f'{fk}__date', f'{fk}_id', 'position').iterator():
        line.parent = getattr(line, fk)
        yield line
