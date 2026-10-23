"""Smoke tests for the PDF / Excel renderers (core.pdf, core.excel). No database needed.

    python manage.py test core.test_render
"""
import re
import shutil
import tempfile
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from django.test import SimpleTestCase

from .excel import BALANCE_FORMAT, DATE_FORMAT, MONEY_FORMAT, render_report_xlsx
from .money import amount_in_words, fmt_money, fmt_qty
from .pdf import format_value, render_invoice_pdf, render_receipt_pdf, render_report_pdf, rich

RUPEE = chr(0x20B9)
GUJ_NAME = 'ઘનશ્યામભાઈ નાના'


def page_count(pdf):
    return len(re.findall(rb'/Type /Page\b', pdf))


def media_box(pdf):
    match = re.search(rb'/MediaBox \[\s*0 0 ([\d.]+) ([\d.]+)', pdf)
    return float(match.group(1)), float(match.group(2))


def make_business(logo_path=None):
    return {
        'name': 'Shreeji Solar & Electricals', 'address': 'Shop 12, Kalavad Road\nRajkot 360005',
        'phone': '+91 98250 12345', 'email': 'accounts@shreejisolar.in', 'gstin': '24ABCDE1234F1Z5',
        'logo_path': logo_path,
    }


def make_items(count):
    items = []
    for i in range(1, count + 1):
        amount = Decimal(500 + i * 37)
        items.append({
            'sno': i, 'name': 'પેનલ 90W' if i % 2 else f'Inverter 1KVA model {i}',
            'description': 'ક્ષમતા: 90W, (શ્રી ઘનશ્યામભાઈ)-12V' if i % 3 == 0 else '',
            'hsn': '85414300', 'qty': fmt_qty(Decimal(i % 4 + 1)), 'unit': 'નંગ' if i % 2 else 'Nos',
            'rate': fmt_money(amount, symbol=''), 'discount': '10.00' if i % 5 == 0 else '0.00',
            'discount_label': '10%' if i % 5 == 0 else '', 'tax_label': '18%', 'tax': '90.00',
            'amount': fmt_money(amount * (i % 4 + 1), symbol=''),
        })
    return items


def make_invoice(items=60, logo_path=None, **overrides):
    doc = {
        'doc_type': 'sale', 'title': 'TAX INVOICE', 'business': make_business(logo_path),
        'party_label': 'Bill To',
        'party': {'name': GUJ_NAME, 'phone': '9876543210', 'address': 'મોરબી રોડ,\nRajkot', 'gstin': '24AAACG1234K1Z9'},
        'meta': [('Invoice No.', 'Bil.12'), ('Date', '16/09/2026'), ('Due Date', '30/09/2026'), ('Payment', 'Credit')],
        'status': 'partial', 'status_label': 'PARTIALLY PAID',
        'columns': {'hsn': True, 'discount': True, 'tax': True},
        'items': make_items(items), 'total_qty': '150',
        'summary': [{'label': 'Subtotal', 'value': RUPEE + '9,000.00'}, {'label': 'Discount', 'value': '-' + RUPEE + '100.00'},
                    {'label': 'Tax', 'value': RUPEE + '1,100.00'}, {'label': 'Round Off', 'value': RUPEE + '0.00'}],
        'total': RUPEE + '10,000.00', 'paid_label': 'Received', 'paid': RUPEE + '5,000.00',
        'balance_label': 'Balance Due', 'balance': RUPEE + '5,000.00',
        'amount_in_words': amount_in_words(10000),
        'payments': [{'number': 'Receipt 3', 'date': '16/09/2026', 'mode': 'HDFC બેંક', 'amount': RUPEE + '5,000.00'}],
        'notes': 'માલ મળ્યા પછી 7 દિવસમાં ચુકવણી. Delivery -> Rajkot.',
        'terms': '1. Goods once sold will not be taken back.\n2. વ્યાજ 18% p.a. લાગુ થશે.',
        'bank': {'bank_name': 'HDFC Bank', 'account_name': 'Shreeji Solar', 'account_number': '50200012345678',
                 'ifsc': 'HDFC0001234', 'branch': 'Rajkot', 'upi_id': 'shreeji@hdfcbank'},
        'show_signature': True, 'footer_note': 'Thank you for your business! આભાર.',
        'tax_summary': [{'rate': 'GST 18%', 'taxable': '9,000.00', 'tax': '1,100.00'}],
    }
    doc.update(overrides)
    return doc


def make_minimal_invoice():
    return {
        'doc_type': 'purchase', 'title': 'PURCHASE BILL',
        'business': {'name': 'Tallin Traders', 'address': '', 'phone': '', 'email': '', 'gstin': '',
                     'logo_path': '/definitely/missing/logo.png'},
        'party_label': 'Supplier', 'party': {'name': '', 'phone': '', 'address': '', 'gstin': ''},
        'meta': [('Bill No.', '45')], 'status': 'unpaid', 'status_label': 'UNPAID',
        'columns': {'hsn': False, 'discount': False, 'tax': False}, 'items': [], 'total_qty': '',
        'summary': [], 'total': RUPEE + '10,000.00', 'paid_label': '', 'paid': '',
        'balance_label': 'Balance Due', 'balance': RUPEE + '10,000.00', 'amount_in_words': '',
        'payments': [], 'notes': '', 'terms': '', 'bank': None, 'show_signature': False,
        'footer_note': '', 'tax_summary': [],
    }


def make_receipt(logo_path=None, allocations=3):
    return {
        'doc_type': 'payment_in', 'title': 'PAYMENT RECEIPT', 'business': make_business(logo_path),
        'party_label': 'Received From', 'party': {'name': GUJ_NAME, 'phone': '9876543210', 'address': 'Rajkot', 'gstin': ''},
        'meta': [('Receipt No.', 'Receipt 3'), ('Date', '16/09/2026'), ('Payment Mode', 'HDFC Bank'), ('Reference No.', 'UTR123')],
        'amount': RUPEE + '5,000.00', 'amount_in_words': amount_in_words(5000),
        'allocations': [{'number': f'Bil.{i}', 'date': '01/09/2026', 'bill_total': RUPEE + '3,000.00',
                         'amount': RUPEE + '1,000.00'} for i in range(allocations)],
        'unallocated_label': 'Advance / Unlinked', 'unallocated': RUPEE + '2,000.00',
        'party_balance_label': 'Current Balance', 'party_balance': RUPEE + '5,000.00 (Receivable)',
        'notes': 'ચેક નં. 004512', 'show_signature': True, 'footer_note': 'Thank you!',
    }


STATEMENT_COLUMNS = [
    {'key': 'date', 'label': 'Date', 'type': 'date'},
    {'key': 'type', 'label': 'Type', 'type': 'text'},
    {'key': 'description', 'label': 'Description', 'type': 'multiline'},
    {'key': 'debit', 'label': 'Debit', 'type': 'money'},
    {'key': 'credit', 'label': 'Credit', 'type': 'money'},
    {'key': 'qty', 'label': 'Qty', 'type': 'qty', 'width': 0.7},
    {'key': 'count', 'label': 'Items', 'type': 'number'},
    {'key': 'balance', 'label': 'Balance', 'type': 'balance'},
]


def make_report(rows=25, orientation='portrait', **overrides):
    data, balance = [], Decimal('0')
    for i in range(rows):
        debit = Decimal(1000 + i) if i % 2 == 0 else None
        credit = None if debit else Decimal('750.50')
        balance += (debit or 0) - (credit or 0)
        data.append({'date': date(2026, 9, 1) + timedelta(days=i % 30), 'type': 'Sale' if debit else 'Payment In',
                     'description': f'{GUJ_NAME}\nBill {i}' if i % 3 == 0 else 'Inverter 1KVA',
                     'debit': debit, 'credit': credit, 'qty': Decimal('2.500'), 'count': i, 'balance': balance})
    report = {
        'title': 'Customer Statement', 'subtitle': 'Period: 01/09/2026 to 16/09/2026', 'business': make_business(),
        'generated_at': '16/09/2026 05:42 PM', 'orientation': orientation,
        'info': [('Party', GUJ_NAME), ('Mobile', '9876543210'), ('Address', 'Rajkot')],
        'summary': [{'label': 'Opening Balance', 'value': Decimal('0'), 'type': 'balance'},
                    {'label': 'Total Debit', 'value': Decimal('123456.5'), 'type': 'money'},
                    {'label': 'Closing Balance', 'value': balance, 'type': 'balance'},
                    {'label': 'Qty', 'value': Decimal('3.000'), 'type': 'qty'},
                    {'label': 'Top Party', 'value': GUJ_NAME, 'type': 'text'}],
        'columns': STATEMENT_COLUMNS, 'rows': data,
        'totals': {'debit': Decimal('5000'), 'credit': Decimal('2500'), 'balance': balance},
        'notes': ['Closing Balance: ' + RUPEE + '5,000.00 (Receivable)'],
    }
    report.update(overrides)
    return report


class RenderTestBase(SimpleTestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        from PIL import Image, ImageDraw
        cls.tmpdir = Path(tempfile.mkdtemp(prefix='tallin-render-'))
        cls.logo = cls.tmpdir / 'logo.png'
        image = Image.new('RGBA', (300, 200), (0, 0, 0, 0))
        ImageDraw.Draw(image).rounded_rectangle((5, 5, 295, 195), radius=30, fill=(29, 78, 216, 255))
        image.save(cls.logo)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmpdir, ignore_errors=True)
        super().tearDownClass()

    def assertPdf(self, data):
        self.assertIsInstance(data, bytes)
        self.assertTrue(data.startswith(b'%PDF'), data[:10])
        self.assertTrue(data.rstrip().endswith(b'%%EOF'))


class RichTextTests(SimpleTestCase):
    def test_escapes_markup(self):
        self.assertEqual(rich('<b>A & B</b>'), '&lt;b&gt;A &amp; B&lt;/b&gt;')

    def test_wraps_gujarati_runs_with_inner_spaces(self):
        markup = rich(f'Customer {GUJ_NAME} owes {RUPEE}5,000')
        self.assertIn(f'<font name="NotoSansGujarati">{GUJ_NAME}</font>', markup)
        self.assertTrue(markup.startswith('Customer '))
        self.assertIn('owes ' + RUPEE + '5,000', markup)

    def test_bold_newlines_and_fallbacks(self):
        markup = rich('પેનલ\nA ' + chr(0x2192) + ' B' + chr(0x1F600), bold=True)
        self.assertIn('<font name="NotoSansGujarati-Bold">પેનલ</font><br/>', markup)
        self.assertIn('A -&gt; B', markup)
        self.assertNotIn(chr(0x1F600), markup)

    def test_empty_values(self):
        self.assertEqual(rich(None), '')
        self.assertEqual(rich(''), '')

    def test_format_value(self):
        self.assertEqual(format_value(Decimal('123456'), 'money'), '1,23,456.00')
        self.assertEqual(format_value(Decimal('-1234.5'), 'money'), '-1,234.50')
        self.assertEqual(format_value(Decimal('123456'), 'money', symbol=True), RUPEE + '1,23,456.00')
        self.assertEqual(format_value(Decimal('1234'), 'balance'), '1,234.00 Dr')
        self.assertEqual(format_value(Decimal('-1234'), 'balance'), '1,234.00 Cr')
        self.assertEqual(format_value(Decimal('0'), 'balance'), '0.00')
        self.assertEqual(format_value(Decimal('2.500'), 'qty'), '2.5')
        self.assertEqual(format_value(Decimal('3.000'), 'qty'), '3')
        self.assertEqual(format_value(7, 'number'), '7')
        self.assertEqual(format_value(date(2026, 9, 16), 'date'), '16/09/2026')
        self.assertEqual(format_value(None, 'money'), '')


class InvoicePdfTests(RenderTestBase):
    def test_all_formats_long_gujarati_invoice(self):
        doc = make_invoice(60, str(self.logo))
        for fmt in ('a4', 'a5', 'thermal80', 'thermal58'):
            with self.subTest(fmt=fmt):
                pdf = render_invoice_pdf(doc, fmt)
                self.assertPdf(pdf)
                if fmt.startswith('thermal'):
                    width, height = media_box(pdf)
                    self.assertEqual(page_count(pdf), 1)
                    self.assertAlmostEqual(width, (80 if fmt == 'thermal80' else 58) * 72 / 25.4, places=1)
                    self.assertGreater(height, width)
                else:
                    self.assertGreaterEqual(page_count(pdf), 3)

    def test_a5_short_invoice(self):
        pdf = render_invoice_pdf(make_invoice(5, str(self.logo), status='paid', status_label='PAID'), 'a5')
        self.assertPdf(pdf)
        self.assertAlmostEqual(media_box(pdf)[0], 419.53, places=1)

    def test_zero_items_and_empty_optional_fields(self):
        for fmt in ('a4', 'a5', 'thermal80', 'thermal58'):
            with self.subTest(fmt=fmt):
                pdf = render_invoice_pdf(make_minimal_invoice(), fmt)
                self.assertPdf(pdf)
                self.assertEqual(page_count(pdf), 1)

    def test_unknown_format_falls_back_to_a4(self):
        pdf = render_invoice_pdf(make_invoice(3), 'letter')
        self.assertPdf(pdf)
        self.assertAlmostEqual(media_box(pdf)[0], 595.28, places=1)

    def test_mixed_script_words_and_hostile_text(self):
        doc = make_invoice(4, '/missing/logo.png')
        doc['party']['name'] = f'<b>{GUJ_NAME},</b> & (પેનલ)-90W 5000ની ' + chr(0x1F64F)
        doc['items'][0]['name'] = 'ક્ષ' * 40
        doc['notes'] = 'Line1\r\nLine2\x07'
        for fmt in ('a4', 'thermal58'):
            with self.subTest(fmt=fmt):
                self.assertPdf(render_invoice_pdf(doc, fmt))

    def test_missing_keys(self):
        for fmt in ('a4', 'thermal80'):
            with self.subTest(fmt=fmt):
                self.assertPdf(render_invoice_pdf({}, fmt))


class ReceiptPdfTests(RenderTestBase):
    def test_all_formats(self):
        for fmt in ('a4', 'a5', 'thermal80', 'thermal58'):
            with self.subTest(fmt=fmt):
                pdf = render_receipt_pdf(make_receipt(str(self.logo)), fmt)
                self.assertPdf(pdf)
                self.assertEqual(page_count(pdf), 1)

    def test_payment_out_without_allocations(self):
        doc = make_receipt(allocations=0)
        doc.update({'doc_type': 'payment_out', 'title': 'PAYMENT VOUCHER', 'party_label': 'Paid To',
                    'unallocated_label': '', 'unallocated': '', 'notes': '', 'show_signature': False})
        for fmt in ('a4', 'thermal58'):
            with self.subTest(fmt=fmt):
                self.assertPdf(render_receipt_pdf(doc, fmt))


class ReportPdfTests(RenderTestBase):
    def test_statement_portrait(self):
        pdf = render_report_pdf(make_report(25))
        self.assertPdf(pdf)
        width, height = media_box(pdf)
        self.assertLess(width, height)

    def test_landscape_many_rows(self):
        pdf = render_report_pdf(make_report(250, 'landscape', info=[]))
        self.assertPdf(pdf)
        width, height = media_box(pdf)
        self.assertGreater(width, height)
        self.assertGreater(page_count(pdf), 3)

    def test_zero_rows(self):
        pdf = render_report_pdf(make_report(0, totals=None, summary=[], notes=[]))
        self.assertPdf(pdf)
        self.assertEqual(page_count(pdf), 1)

    def test_missing_keys(self):
        self.assertPdf(render_report_pdf({}))


class ReportXlsxTests(RenderTestBase):
    def load(self, data):
        from openpyxl import load_workbook
        self.assertTrue(data.startswith(b'PK'))
        return load_workbook(BytesIO(data)).active

    def header_row(self, ws):
        for row in range(1, ws.max_row + 1):
            if ws.cell(row, 1).value == 'Date' and ws.cell(row, 2).value == 'Type':
                return row
        self.fail('header row not found')

    def test_statement_types_and_layout(self):
        report = make_report(10)
        report['rows'][1]['description'] = '=HYPERLINK("http://example.com")'
        ws = self.load(render_report_xlsx(report))
        header = self.header_row(ws)
        self.assertEqual(ws.cell(1, 1).value, 'Shreeji Solar & Electricals')
        self.assertTrue(ws.cell(1, 1).font.b)
        self.assertEqual(ws.freeze_panes, f'A{header + 1}')
        self.assertEqual(ws.auto_filter.ref, f'A{header}:H{header + 10}')
        first = header + 1
        self.assertIsInstance(ws.cell(first, 1).value, datetime)
        self.assertEqual(ws.cell(first, 1).number_format, DATE_FORMAT)
        self.assertEqual(ws.cell(first, 4).number_format, MONEY_FORMAT)
        self.assertEqual(ws.cell(first, 4).value, 1000)
        self.assertEqual(ws.cell(first, 6).value, 2.5)
        self.assertEqual(ws.cell(first, 8).number_format, BALANCE_FORMAT)
        self.assertTrue(ws.cell(first, 3).alignment.wrap_text)
        self.assertEqual(ws.cell(first + 1, 3).data_type, 's')  # never a formula
        totals_row = header + 11
        self.assertEqual(ws.cell(totals_row, 1).value, 'Total')
        self.assertTrue(ws.cell(totals_row, 4).font.b)
        self.assertEqual(ws.page_setup.orientation, 'portrait')
        self.assertEqual(ws.page_setup.fitToWidth, 1)

    def test_landscape_zero_rows_and_title(self):
        report = make_report(0, 'landscape', title='Sales: Report [A/B] * very long title here', totals=None)
        ws = self.load(render_report_xlsx(report))
        self.assertLessEqual(len(ws.title), 31)
        self.assertFalse(any(ch in ws.title for ch in '[]:*?/\\'))
        header = self.header_row(ws)
        self.assertEqual(ws.cell(header + 1, 1).value, 'No records for the selected period.')
        self.assertEqual(ws.page_setup.orientation, 'landscape')

    def test_missing_keys(self):
        self.load(render_report_xlsx({}))
