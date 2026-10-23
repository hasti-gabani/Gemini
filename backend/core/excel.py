"""Excel (.xlsx) export for table reports and party/bank statements (openpyxl).

Public API::

    render_report_xlsx(report: dict) -> bytes

Takes the same ``report`` dict as :func:`core.pdf.render_report_pdf`. Values are
written as native Excel types (numbers, real dates) with display formats, so
sorting, filtering and formulas keep working. Importable without Django settings.
"""
import math
from copy import copy
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

from .money import D, fmt_money, fmt_qty, q2, q3

__all__ = ['render_report_xlsx', 'MONEY_FORMAT', 'BALANCE_FORMAT', 'QTY_FORMAT', 'NUMBER_FORMAT', 'DATE_FORMAT']

# Indian digit grouping (12,34,567.00); negatives fall into the last section (-1,234.00).
MONEY_FORMAT = '[>=10000000]##\\,##\\,##\\,##0.00;[>=100000]##\\,##\\,##0.00;##,##0.00'
BALANCE_FORMAT = '#,##0.00 "Dr";#,##0.00 "Cr";0.00'
QTY_FORMAT = 'General'
NUMBER_FORMAT = '0'
DATE_FORMAT = 'dd/mm/yyyy'

NO_RECORDS = 'No records for the selected period.'
NUMERIC_TYPES = ('money', 'balance', 'qty', 'number')
MAX_CELL_CHARS = 32767
WIDTH_CAPS = {'text': 40, 'multiline': 50, 'date': 12, 'money': 18, 'balance': 20, 'qty': 12, 'number': 12}

BRAND = '1D4ED8'
BORDER_COLOR = 'E2E8F0'
THIN = Side(style='thin', color=BORDER_COLOR)
HEADER_SIDE = Side(style='thin', color='1E40AF')
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_BORDER = Border(left=HEADER_SIDE, right=HEADER_SIDE, top=HEADER_SIDE, bottom=HEADER_SIDE)
TOTAL_BORDER = Border(left=THIN, right=THIN, top=Side(style='medium', color=BRAND), bottom=Side(style='thin', color=BRAND))
HEADER_FILL = PatternFill('solid', start_color=BRAND, end_color=BRAND)
TOTAL_FILL = PatternFill('solid', start_color='EFF6FF', end_color='EFF6FF')
HEADER_FONT = Font(bold=True, color='FFFFFF')
BOLD = Font(bold=True)
TONE_FONTS = {'pos': Font(bold=True, color='15803D'), 'neg': Font(bold=True, color='B91C1C')}
MUTED_FONT = Font(italic=True, color='475569')
TOP_LEFT = Alignment(vertical='top')
TOP_WRAP = Alignment(vertical='top', wrap_text=True)
HEADER_LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
HEADER_RIGHT = Alignment(horizontal='right', vertical='center', wrap_text=True)


def _text(value):
    """Excel-safe string: no control characters, CR/LF normalised, length-limited."""
    if value is None:
        return ''
    text = str(value).replace('\r\n', '\n').replace('\r', '\n')
    return ILLEGAL_CHARACTERS_RE.sub('', text)[:MAX_CELL_CHARS]


def _decimal(value):
    if value is None or value == '' or isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        return value
    try:
        return D(value)
    except (ValueError, TypeError, ArithmeticError):
        return None


def _sheet_title(title):
    text = _text(title)
    for ch in '[]:*?/' + chr(92):
        text = text.replace(ch, ' ')
    text = ' '.join(text.split())[:31].strip().strip("'").strip()
    return text or 'Report'


def _excel_value(value, value_type):
    """(native value, number format, display text) for a report value."""
    if value is None or value == '':
        return None, None, ''
    value_type = (value_type or 'text').lower()
    if value_type in ('money', 'balance'):
        amount = _decimal(value)
        if amount is None:
            text = _text(value)
            return text, None, text
        amount = q2(amount)
        display = fmt_money(abs(amount) if value_type == 'balance' else amount, symbol='')
        if value_type == 'balance' and amount != 0:
            display += ' Dr' if amount > 0 else ' Cr'
        return float(amount), MONEY_FORMAT if value_type == 'money' else BALANCE_FORMAT, display
    if value_type == 'qty':
        amount = _decimal(value)
        if amount is None:
            text = _text(value)
            return text, None, text
        number = float(q3(amount))
        return (int(number) if number.is_integer() else number), QTY_FORMAT, fmt_qty(amount)
    if value_type == 'number':
        try:
            number = int(value)
        except (TypeError, ValueError, ArithmeticError):
            text = _text(value)
            return text, None, text
        return number, NUMBER_FORMAT, str(number)
    if value_type == 'date':
        if isinstance(value, datetime):
            value = value.date()
        if isinstance(value, date):
            return value, DATE_FORMAT, '00/00/0000'
        text = _text(value)
        return text, None, text
    text = _text(value)
    return text, None, text


_STYLE_ATTRS = ('number_format', 'font', 'fill', 'border', 'alignment')


def _put(ws, row, column, value, number_format=None, font=None, fill=None, border=None, alignment=None,
         style_cache=None):
    """Write a cell. ``style_cache`` (dict) re-uses the computed style array for identical styles,
    avoiding openpyxl's per-assignment style hashing on large sheets."""
    cell = ws.cell(row=row, column=column)
    cell.value = value
    if isinstance(value, str) and value.startswith('='):
        cell.data_type = 's'  # user text is never evaluated as a formula
    styles = (number_format, font, fill, border, alignment)
    key = tuple(id(s) if s is not None and not isinstance(s, str) else s for s in styles)
    cached = style_cache.get(key) if style_cache is not None else None
    if cached is not None:
        cell._style = copy(cached)
        return cell
    for name, style in zip(_STYLE_ATTRS, styles):
        if style is not None and style != '':
            setattr(cell, name, style)
    if style_cache is not None and hasattr(cell, '_style'):
        style_cache[key] = copy(cell._style)
    return cell


def _pairs(entries):
    out = []
    for entry in entries or []:
        if isinstance(entry, dict):
            label, value = entry.get('label'), entry.get('value')
        else:
            entry = list(entry) + ['', '']
            label, value = entry[0], entry[1]
        if _text(label).strip() or _text(value).strip():
            out.append((_text(label).strip(), value))
    return out


def _line_count(text, width):
    if not text:
        return 1
    usable = max(width - 1, 1)
    return sum(max(1, math.ceil(len(line) / usable)) for line in text.split('\n'))


def render_report_xlsx(report: dict) -> bytes:
    """Report / statement as .xlsx bytes."""
    report = report if isinstance(report, dict) else {}
    columns = [c for c in (report.get('columns') or []) if isinstance(c, dict)]
    rows = [r for r in (report.get('rows') or []) if isinstance(r, dict)]
    totals = report.get('totals') if isinstance(report.get('totals'), dict) else None
    ncols = max(len(columns), 1)
    keys = [c.get('key') for c in columns]
    types = [(_text(c.get('type')).strip() or 'text').lower() for c in columns]
    labels = [_text(c.get('label')).strip() for c in columns]

    wb = Workbook()
    ws = wb.active
    title = _text(report.get('title')).strip() or 'Report'
    ws.title = _sheet_title(title)
    business = report.get('business') or {}
    wb.properties.title = title
    wb.properties.creator = 'Gemini'
    if _text(business.get('name')).strip():
        wb.properties.company = _text(business.get('name')).strip()

    widths = [8.0] * ncols
    row = 1
    if _text(business.get('name')).strip():
        _put(ws, row, 1, _text(business.get('name')).strip(), font=Font(bold=True, size=14))
        ws.row_dimensions[row].height = 20
        row += 1
    _put(ws, row, 1, title, font=Font(bold=True, size=12))
    ws.row_dimensions[row].height = 17
    row += 1
    if _text(report.get('subtitle')).strip():
        _put(ws, row, 1, _text(report.get('subtitle')).strip())
        row += 1
    if _text(report.get('generated_at')).strip():
        _put(ws, row, 1, 'Generated on: ' + _text(report.get('generated_at')).strip(), font=MUTED_FONT)
        row += 1

    label_width = 0
    info = _pairs(report.get('info'))
    if info:
        row += 1
        for label, value in info:
            _put(ws, row, 1, label, font=BOLD)
            parts = (part.strip().rstrip(',').strip() for part in _text(value).split('\n'))
            _put(ws, row, 2, ', '.join(part for part in parts if part))
            label_width = max(label_width, len(label))
            row += 1

    summary = [entry for entry in (report.get('summary') or []) if isinstance(entry, dict)]
    if summary:
        row += 1
        for entry in summary:
            label = _text(entry.get('label')).strip()
            value, number_format, display = _excel_value(entry.get('value'), entry.get('type'))
            _put(ws, row, 1, label, font=BOLD)
            _put(ws, row, 2, value, number_format=number_format, font=TONE_FONTS.get(entry.get('tone'), BOLD),
                 alignment=Alignment(horizontal='right') if isinstance(value, (int, float)) else None)
            label_width = max(label_width, len(label))
            if ncols > 1:
                widths[1] = max(widths[1], len(display) + 2)
            row += 1
    widths[0] = max(widths[0], min(label_width + 2, 28))

    row += 1
    header_row = row
    for i, label in enumerate(labels):
        _put(ws, row, i + 1, label, font=HEADER_FONT, fill=HEADER_FILL, border=HEADER_BORDER,
             alignment=HEADER_RIGHT if types[i] in NUMERIC_TYPES else HEADER_LEFT)
        longest_word = max((len(word) for word in label.split()), default=0)
        widths[i] = max(widths[i], min(max(len(label), longest_word) + 2, 24))

    wrap_columns = [t == 'multiline' for t in types]
    style_cache = {}  # only module-level style constants are passed with this cache
    for data in rows:
        row += 1
        lines = 1
        for i, key in enumerate(keys):
            value, number_format, display = _excel_value(data.get(key), types[i])
            _put(ws, row, i + 1, value, number_format=number_format, border=CELL_BORDER,
                 alignment=TOP_WRAP if wrap_columns[i] else TOP_LEFT, style_cache=style_cache)
            if display:
                longest_line = max(len(line) for line in display.split('\n'))
                cap = WIDTH_CAPS.get(types[i], 40)
                widths[i] = max(widths[i], min(longest_line + 2, cap))
                if wrap_columns[i]:
                    lines = max(lines, _line_count(display, cap))
        if lines > 1:
            ws.row_dimensions[row].height = min(15 * lines, 409)
    last_data_row = row

    if not rows and columns:
        row += 1
        _put(ws, row, 1, NO_RECORDS, font=MUTED_FONT, alignment=Alignment(horizontal='center'))
        if ncols > 1:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)

    if totals is not None and columns:
        row += 1
        for i, key in enumerate(keys):
            value, number_format, display = _excel_value(totals.get(key), types[i])
            if i == 0 and value in (None, ''):
                value, number_format, display = 'Total', None, 'Total'
            _put(ws, row, i + 1, value, number_format=number_format, font=BOLD, fill=TOTAL_FILL,
                 border=TOTAL_BORDER, alignment=TOP_WRAP if wrap_columns[i] else TOP_LEFT)
            if display:
                widths[i] = max(widths[i], min(len(display) + 2, WIDTH_CAPS.get(types[i], 40)))

    notes = [_text(note).strip() for note in (report.get('notes') or []) if _text(note).strip()]
    if notes:
        row += 1
        for note in notes:
            row += 1
            _put(ws, row, 1, note)

    # header row height: enough lines for the longest wrapped label
    header_lines = max((_line_count(label, widths[i]) for i, label in enumerate(labels)), default=1)
    ws.row_dimensions[header_row].height = max(20, 15 * header_lines)
    for i, width in enumerate(widths):
        ws.column_dimensions[get_column_letter(i + 1)].width = round(width, 1)

    last_col = get_column_letter(ncols)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if columns:
        ws.auto_filter.ref = f'A{header_row}:{last_col}{max(last_data_row, header_row)}'
        ws.print_title_rows = f'{header_row}:{header_row}'

    landscape = _text(report.get('orientation')).strip().lower() == 'landscape'
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.6, header=0.3, footer=0.3)
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = 'Page &P of &N'

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
